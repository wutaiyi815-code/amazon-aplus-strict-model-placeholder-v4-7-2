"""Shared, full-range Excel reader for Amazon A+ scripts.

Use normal openpyxl mode everywhere. Read-only streaming mode trusts worksheet
dimension metadata and can silently stop at a stale declared range such as A1.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any


EXCEL_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
DEFAULT_EXCEL_NAME = "1-产品属性表.xlsx"


def find_product_workbook(product_dir: Path) -> Path | None:
    preferred = product_dir / DEFAULT_EXCEL_NAME
    if preferred.exists():
        return preferred
    candidates = sorted(
        [
            path
            for path in product_dir.iterdir()
            if path.is_file()
            and path.suffix.lower() in EXCEL_EXTENSIONS
            and not path.name.startswith("~$")
        ],
        key=lambda item: item.name.casefold(),
    )
    return candidates[0] if candidates else None


def read_active_worksheet_rows(path: Path) -> tuple[str, list[tuple[Any, ...]], str]:
    """Return all physical rows from the active sheet without trusting stale dimensions."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=False)
    try:
        sheet = workbook.active
        rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
        return sheet.title, rows, sheet.calculate_dimension()
    finally:
        workbook.close()
