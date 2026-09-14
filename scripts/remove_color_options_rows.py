#!/usr/bin/env python
"""Remove color_options rows from every SKU attribute workbook before A+ work."""

from __future__ import annotations

import argparse
import json
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


EXCEL_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
EXCLUDED_DIR_NAMES = {"logo"}
EXCLUDED_DIR_PREFIXES = ("_aplus", "输出结果")
TARGET_KEY = "color_options"


def is_product_dir(path: Path) -> bool:
    lowered = path.name.casefold()
    return (
        path.is_dir()
        and lowered not in EXCLUDED_DIR_NAMES
        and not any(lowered.startswith(prefix.casefold()) for prefix in EXCLUDED_DIR_PREFIXES)
    )


def find_attribute_workbook(product_dir: Path) -> Path | None:
    preferred = product_dir / "1-产品属性表.xlsx"
    if preferred.exists():
        return preferred
    candidates = sorted(
        (
            path
            for path in product_dir.iterdir()
            if path.is_file()
            and path.suffix.casefold() in EXCEL_EXTENSIONS
            and not path.name.startswith("~$")
        ),
        key=lambda item: item.name.casefold(),
    )
    return candidates[0] if candidates else None


def normalized_key(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip().rstrip(":：").casefold()


def matching_rows(sheet: Any) -> list[dict[str, Any]]:
    matches: list[dict[str, Any]] = []
    for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        values = list(row)
        first = next((value for value in values if value is not None and str(value).strip()), None)
        if normalized_key(first) == TARGET_KEY:
            matches.append(
                {
                    "row": row_index,
                    "values": [None if value is None else str(value) for value in values],
                }
            )
    return matches


def workbook_matches(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=False, read_only=False, keep_vba=path.suffix.casefold() in {".xlsm", ".xltm"})
    try:
        found: list[dict[str, Any]] = []
        for sheet in workbook.worksheets:
            for match in matching_rows(sheet):
                found.append({"sheet": sheet.title, **match})
        return found
    finally:
        workbook.close()


def remove_from_workbook(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=False, read_only=False, keep_vba=path.suffix.casefold() in {".xlsm", ".xltm"})
    removed: list[dict[str, Any]] = []
    try:
        for sheet in workbook.worksheets:
            matches = matching_rows(sheet)
            for match in sorted(matches, key=lambda item: int(item["row"]), reverse=True):
                sheet.delete_rows(int(match["row"]), 1)
            removed.extend({"sheet": sheet.title, **match} for match in matches)
        if removed:
            workbook.save(path)
    finally:
        workbook.close()
    return removed


def process_root(root: Path, *, dry_run: bool = False, audit_path: Path | None = None) -> dict[str, Any]:
    root = root.expanduser().resolve()
    if not root.exists() or not root.is_dir():
        raise SystemExit(f"Target root does not exist or is not a directory: {root}")

    run_stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    backup_root = root / "_aplus_color_options_backups" / run_stamp
    audit_path = audit_path or root / "_aplus_color_options_removal.json"
    records: list[dict[str, Any]] = []
    errors: list[str] = []

    for product_dir in sorted((path for path in root.iterdir() if is_product_dir(path)), key=lambda item: item.name.casefold()):
        workbook_path = find_attribute_workbook(product_dir)
        if workbook_path is None:
            records.append({"sku": product_dir.name, "status": "no_attribute_workbook", "removed_rows": []})
            continue
        record: dict[str, Any] = {
            "sku": product_dir.name,
            "workbook": str(workbook_path),
            "status": "pending",
            "removed_rows": [],
        }
        try:
            planned = workbook_matches(workbook_path)
            record["planned_rows"] = planned
            if not planned:
                record["status"] = "no_color_options_row"
            elif dry_run:
                record["status"] = "dry_run"
            else:
                backup_path = backup_root / product_dir.name / workbook_path.name
                backup_path.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(workbook_path, backup_path)
                record["backup"] = str(backup_path)
                removed = remove_from_workbook(workbook_path)
                remaining = workbook_matches(workbook_path)
                if remaining:
                    raise RuntimeError(f"Verification failed; {len(remaining)} color_options row(s) remain")
                record["removed_rows"] = removed
                record["status"] = "removed"
        except Exception as exc:  # noqa: BLE001 - aggregate all workbook failures in the audit
            record["status"] = "error"
            record["error"] = str(exc)
            errors.append(f"{product_dir.name}: {exc}")
        records.append(record)

    payload = {
        "root": str(root),
        "timestamp_utc": datetime.now(timezone.utc).isoformat(),
        "dry_run": dry_run,
        "target_key": TARGET_KEY,
        "backup_root": str(backup_root) if not dry_run else None,
        "products_checked": len(records),
        "workbooks_changed": sum(1 for item in records if item["status"] == "removed"),
        "rows_removed": sum(len(item.get("removed_rows", [])) for item in records),
        "errors": errors,
        "records": records,
    }
    audit_path.parent.mkdir(parents=True, exist_ok=True)
    audit_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return payload


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", required=True, help="A+ target root whose direct children are SKU folders.")
    parser.add_argument("--audit", help="Optional audit JSON path.")
    parser.add_argument("--dry-run", action="store_true", help="Report matching rows without changing workbooks.")
    args = parser.parse_args()

    payload = process_root(
        Path(args.root),
        dry_run=args.dry_run,
        audit_path=Path(args.audit).expanduser().resolve() if args.audit else None,
    )
    print(json.dumps({key: payload[key] for key in ("products_checked", "workbooks_changed", "rows_removed", "errors")}, ensure_ascii=False))
    return 1 if payload["errors"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
