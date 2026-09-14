from __future__ import annotations

import argparse
import json
import math
import os
import re
from pathlib import Path
from typing import Any

from PIL import Image, ImageOps, ImageDraw
from openpyxl import load_workbook


IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}
EXCEL_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm"}


def find_excel(product_dir: Path) -> Path | None:
    preferred = product_dir / "1-äº§å“å±žæ€§è¡¨.xlsx"
    if preferred.exists():
        return preferred
    candidates = sorted(
        [
            path
            for path in product_dir.iterdir()
            if path.is_file()
            and path.suffix.lower() in EXCEL_EXTENSIONS
            and not path.name.startswith("~$")
        ],
        key=lambda item: item.name.casefold(),
    )
    return candidates[0] if candidates else None


def read_excel_facts(path: Path) -> dict[str, str]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    facts: dict[str, str] = {}
    for row in ws.iter_rows(values_only=True):
        cells = [str(v).strip() if v is not None else "" for v in row]
        if len(cells) >= 2 and cells[0] and cells[0].lower() != "field":
            facts[cells[0]] = cells[1]
    wb.close()
    return facts


def occasions_from_facts(facts: dict[str, str]) -> list[dict[str, str]]:
    occasions = []
    for key, value in facts.items():
        if re.fullmatch(r"occasion_\d+", key.lower()) and value.strip():
            idx = int(key.split("_", 1)[1])
            occasions.append({"index": idx, "name": value.strip()})
    return sorted(occasions, key=lambda item: item["index"])


def images_in(folder: Path) -> list[Path]:
    ignored_dirs = {"_aplus_work", "_aplus_creative_work", "_influencer_selfie_work"}
    out = []
    for path in folder.iterdir():
        if path.is_dir() and path.name in ignored_dirs:
            continue
        if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS and not path.name.startswith("~$"):
            out.append(path)
    return sorted(out, key=lambda p: p.name.lower())


def infer_target_gender(facts: dict[str, str]) -> str:
    text = " ".join(
        facts.get(k, "")
        for k in ("target_customer", "product_type", "must_include", "feature_1", "feature_2")
    ).lower()
    if any(word in text for word in ["women", "woman", "female", "girl", "ladies", "womens", "women's"]):
        return "women"
    if any(word in text for word in ["men", "man", "male", "boy", "mens", "men's"]):
        return "men"
    if "unisex" in text:
        return "unisex"
    return "unspecified"


def normalized_model_gender(value: str) -> str:
    gender = str(value or "").strip().casefold().replace("-", "_")
    if gender in {"woman", "women", "female", "female_presenting"}:
        return "women"
    if gender in {"man", "men", "male", "male_presenting"}:
        return "men"
    if gender in {"ambiguous", "androgynous", "unisex"}:
        return "unisex"
    return "unspecified"


def select_tagged_model_identity(product_dir: Path) -> tuple[Path | None, str]:
    """Prefer the clearest usable confirmed-person tag as the gender source of truth."""
    tags_path = product_dir / "_aplus_creative_work" / "aigc_model_view_tags.json"
    if not tags_path.exists():
        return None, "unspecified"
    try:
        raw = json.loads(tags_path.read_text(encoding="utf-8-sig"))
    except Exception:
        return None, "unspecified"
    if not isinstance(raw, dict):
        return None, "unspecified"
    ranked: list[tuple[int, str, Path, str]] = []
    for rel, tag in raw.items():
        if not isinstance(tag, dict):
            continue
        if len(Path(str(rel)).parts) == 1:
            continue
        if not bool(tag.get("contains_person")):
            continue
        if str(tag.get("subject_type") or "").casefold().replace("-", "_") != "model_person":
            continue
        if not bool(tag.get("usable_for_identity")):
            continue
        candidate = (product_dir / str(rel)).resolve()
        if not candidate.exists() or candidate.suffix.casefold() not in IMAGE_EXTENSIONS:
            continue
        view = str(tag.get("view") or "unknown").casefold().replace("-", "_")
        framing = str(tag.get("framing") or "unknown").casefold().replace("-", "_")
        quality = str(tag.get("identity_quality") or "unknown").casefold()
        score = {
            "front": 50,
            "three_quarter": 40,
            "side": 20,
            "back": 0,
        }.get(view, 10)
        score += {"full_body": 20, "half_body": 12, "upper_body": 8}.get(framing, 0)
        score += {"high": 10, "medium": 5}.get(quality, 0)
        ranked.append((score, str(candidate).casefold(), candidate, normalized_model_gender(tag.get("gender_presentation", ""))))
    if not ranked:
        return None, "unspecified"
    ranked.sort(key=lambda item: (-item[0], item[1]))
    _, _, identity, gender = ranked[0]
    return identity, gender


def allowed_root_product_images(product_dir: Path) -> list[Path]:
    """Keep root-level non-person product images, but never root-level model/person images."""
    images = images_in(product_dir)
    tags_path = product_dir / "_aplus_creative_work" / "aigc_model_view_tags.json"
    try:
        tags = json.loads(tags_path.read_text(encoding="utf-8-sig")) if tags_path.exists() else {}
    except Exception:
        tags = {}
    allowed: list[Path] = []
    for path in images:
        rel = path.name
        tag = tags.get(rel) if isinstance(tags, dict) else None
        if isinstance(tag, dict):
            subject = str(tag.get("subject_type") or "").casefold().replace("-", "_")
            if tag.get("contains_person") is True or subject == "model_person":
                continue
        allowed.append(path)
    return allowed


def resolve_target_gender(model_gender: str, facts: dict[str, str]) -> tuple[str, str]:
    normalized = normalized_model_gender(model_gender)
    if normalized != "unspecified":
        return normalized, "model_identity_reference"
    return infer_target_gender(facts), "attribute_table_fallback"


def default_subject(facts: dict[str, str]) -> str:
    product_type = facts.get("product_type", "").lower()
    detail_text = " ".join(facts.get(k, "") for k in ("must_include", "feature_1", "feature_2")).lower()
    text = f"{product_type} {detail_text}"
    if any(word in product_type for word in ["pant", "trouser", "sweatpant", "jean"]):
        return "pants"
    if "skirt" in product_type:
        return "skirt"
    if "dress" in product_type:
        return "dress"
    if any(word in product_type for word in ["shorts", "short pant", "bermuda"]):
        return "shorts"
    if any(word in text for word in ["hoodie", "pullover", "sweater", "tee", "shirt", "top"]):
        return "upper garment"
    return "product garment"


def choose_front_reference(product_images: list[Path]) -> Path | None:
    candidates = [
        p for p in product_images
        if "logo" not in p.stem.lower() and "属性" not in p.name and "preview" not in p.stem.lower()
    ]
    if not candidates:
        return None
    model_like = [p for p in candidates if not p.name.upper().startswith("PP_")]
    search_pool = model_like or candidates
    preferred = ["1.jpg", "1 .jpg", "2.jpg", "3.jpg", "4.jpg", "5.jpg", "7.jpg"]
    by_name = {p.name.lower(): p for p in search_pool}
    for name in preferred:
        if name in by_name:
            return by_name[name]
    for path in search_pool:
        if path.stem.lower().startswith("res_"):
            return path
    return search_pool[0]


def make_contact_sheet(images: list[Path], out_path: Path, title: str, thumb_width: int = 240) -> None:
    if not images:
        return
    thumbs = []
    font_pad = 34
    for image_path in images:
        try:
            img = Image.open(image_path).convert("RGB")
            img.thumbnail((thumb_width, thumb_width), Image.Resampling.LANCZOS)
            canvas = Image.new("RGB", (thumb_width, thumb_width + font_pad), "white")
            x = (thumb_width - img.width) // 2
            canvas.paste(img, (x, 0))
            draw = ImageDraw.Draw(canvas)
            label = image_path.name[:32]
            draw.text((4, thumb_width + 6), label, fill=(0, 0, 0))
            thumbs.append(canvas)
        except Exception:
            continue
    cols = min(4, len(thumbs))
    rows = math.ceil(len(thumbs) / cols)
    sheet = Image.new("RGB", (cols * thumb_width, rows * (thumb_width + font_pad) + 40), "white")
    draw = ImageDraw.Draw(sheet)
    draw.text((8, 8), title, fill=(0, 0, 0))
    y0 = 40
    for i, thumb in enumerate(thumbs):
        x = (i % cols) * thumb_width
        y = y0 + (i // cols) * (thumb_width + font_pad)
        sheet.paste(thumb, (x, y))
    out_path.parent.mkdir(parents=True, exist_ok=True)
    sheet.save(out_path, quality=92)


def main() -> int:
    parser = argparse.ArgumentParser(description="Prepare influencer selfie batch from product folders.")
    parser.add_argument("--root", required=True)
    parser.add_argument("--background-root", default=os.environ.get("APLUS_BACKGROUND_ROOT"), help="Background image folder; defaults to APLUS_BACKGROUND_ROOT.")
    args = parser.parse_args()

    root = Path(args.root).expanduser().resolve()
    if not args.background_root:
        parser.error("Provide --background-root or set APLUS_BACKGROUND_ROOT.")
    background_root = Path(args.background_root).expanduser().resolve()
    if not root.exists():
        raise SystemExit(f"Root does not exist: {root}")
    if not background_root.exists():
        raise SystemExit(f"Background root does not exist: {background_root}")

    background_images = images_in(background_root)
    manifest: dict[str, Any] = {
        "root": str(root),
        "background_root": str(background_root),
        "products": [],
        "background_images": [str(p) for p in background_images],
    }
    make_contact_sheet(background_images, root / "_influencer_background_contact_sheet.jpg", "Influencer background candidates")

    for product_dir in sorted([p for p in root.iterdir() if p.is_dir()]):
        xlsx = find_excel(product_dir)
        if not xlsx:
            continue
        facts = read_excel_facts(xlsx)
        occasions = occasions_from_facts(facts)
        product_images = allowed_root_product_images(product_dir)
        tagged_identity, tagged_gender = select_tagged_model_identity(product_dir)
        work_dir = product_dir / "_influencer_selfie_work"
        work_dir.mkdir(parents=True, exist_ok=True)
        contact_images = list(product_images)
        if tagged_identity and tagged_identity not in contact_images:
            contact_images.append(tagged_identity)
        make_contact_sheet(contact_images, work_dir / "product_contact_sheet.jpg", product_dir.name)
        front_ref = tagged_identity
        target_gender, target_gender_source = resolve_target_gender(tagged_gender, facts)
        plan = {
            "product_id": product_dir.name,
            "product_dir": str(product_dir),
            "excel_path": str(xlsx),
            "facts": facts,
            "target_gender": target_gender,
            "target_gender_source": target_gender_source,
            "target_customer_reference": facts.get("target_customer", ""),
            "product_subject": default_subject(facts),
            "front_model_reference": str(front_ref) if front_ref else "",
            "model_identity_reference": str(front_ref) if front_ref else "",
            "model_identity_notes": "",
            "product_accuracy_notes": "",
            "occasions": [
                {
                    "index": item["index"],
                    "name": item["name"],
                    "background_reference": "",
                    "background_reason": "",
                    "prompt_path": "",
                    "output_path": "",
                }
                for item in occasions
            ],
        }
        plan_path = work_dir / "selfie_plan.json"
        plan_path.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
        manifest["products"].append(str(plan_path))

    manifest_path = root / "_influencer_selfie_manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(manifest_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
